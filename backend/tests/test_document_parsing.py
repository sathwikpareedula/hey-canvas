import zipfile
from pathlib import Path

import pytest
from app.documents.models import DocumentLimits
from app.documents.parsing import (
    DocumentParseError,
    detect_document_type,
    normalize_text,
    parse_document,
)
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


@pytest.fixture
def limits() -> DocumentLimits:
    return DocumentLimits(
        max_source_bytes=5_000_000,
        max_characters=10_000,
        max_units=20,
        max_unit_characters=2_000,
        max_archive_entries=200,
        max_archive_uncompressed_bytes=5_000_000,
        max_archive_ratio=200,
    )


def make_pdf(path: Path, texts: list[str]) -> None:
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


def test_pdf_preserves_page_boundaries_and_numbers(tmp_path: Path, limits: DocumentLimits) -> None:
    path = tmp_path / "lecture.pdf"
    make_pdf(path, ["Inverse kinematics", "Gradient descent"])
    assert detect_document_type(
        path, filename=path.name, content_type="application/pdf", limits=limits
    ) == "pdf"
    result = parse_document(path, "pdf", limits)
    assert result.status == "parsed"
    assert [unit.page_number for unit in result.units] == [1, 2]
    assert "Inverse kinematics" in result.units[0].text
    assert "Gradient descent" in result.units[1].text


def test_empty_pdf_is_needs_ocr(tmp_path: Path, limits: DocumentLimits) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as output:
        writer.write(output)
    result = parse_document(path, "pdf", limits)
    assert result.status == "needs_ocr"
    assert result.units[0].page_number == 1
    assert result.units[0].text == ""


def test_malformed_pdf_is_isolated(tmp_path: Path, limits: DocumentLimits) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-not-valid")
    with pytest.raises(DocumentParseError) as error:
        parse_document(path, "pdf", limits)
    assert error.value.category in {"malformed_document", "parser_error"}


@pytest.mark.parametrize("kind,suffix", [("pptx", ".pptx"), ("docx", ".docx")])
def test_malformed_office_document_is_isolated(
    tmp_path: Path, limits: DocumentLimits, kind: str, suffix: str
) -> None:
    path = tmp_path / f"broken{suffix}"
    path.write_bytes(b"PK\x03\x04broken")
    with pytest.raises(DocumentParseError):
        parse_document(path, kind, limits)  # type: ignore[arg-type]


def test_pptx_extracts_slide_titles_text_and_tables(
    tmp_path: Path, limits: DocumentLimits
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Kinematics"
    slide.placeholders[1].text = "Forward model"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "joint"
    table.cell(0, 1).text = "angle"
    path = tmp_path / "lecture.pptx"
    presentation.save(path)
    result = parse_document(path, "pptx", limits)
    assert len(result.units) == 1
    assert result.units[0].slide_number == 1
    assert result.units[0].heading == "Kinematics"
    assert "Forward model" in result.units[0].text
    assert "joint | angle" in result.units[0].text


def test_docx_preserves_heading_paragraph_table_order(
    tmp_path: Path, limits: DocumentLimits
) -> None:
    document = Document()
    document.add_paragraph("Preface")
    document.add_heading("Control", level=1)
    document.add_paragraph("State feedback")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "gain"
    table.cell(0, 1).text = "value"
    path = tmp_path / "notes.docx"
    document.save(path)
    result = parse_document(path, "docx", limits)
    assert [unit.heading for unit in result.units] == [None, "Control"]
    assert result.units[0].text == "Preface"
    assert result.units[1].text == "State feedback\ngain | value"


@pytest.mark.parametrize(
    "suffix,kind", [(".txt", "text"), (".md", "markdown"), (".csv", "csv")]
)
def test_text_formats_decode_normalize_and_preserve_structure(
    tmp_path: Path, limits: DocumentLimits, suffix: str, kind: str
) -> None:
    path = tmp_path / f"notes{suffix}"
    path.write_bytes(b"alpha\r\n\x00beta\r\n\r\n\r\ngamma")
    result = parse_document(path, kind, limits)  # type: ignore[arg-type]
    assert result.units[0].text == "alpha\nbeta\n\ngamma"


def test_csv_formula_text_is_never_executed(tmp_path: Path, limits: DocumentLimits) -> None:
    path = tmp_path / "grades.csv"
    path.write_text("name,value\nitem,=SUM(1,2)", encoding="utf-8")
    result = parse_document(path, "csv", limits)
    assert "=SUM(1,2)" in result.units[0].text


def test_normalization_is_conservative() -> None:
    assert normalize_text("  x = y  \r\n\x00\r\n\r\n\r\n z ") == "x = y\n\n z"


@pytest.mark.parametrize(
    "name,body,category",
    [
        ("malware.pdf", b"MZnot-a-pdf", "unsupported_executable"),
        ("fake.pdf", b"plain text", "type_mismatch"),
    ],
)
def test_type_mismatch_and_executable_are_rejected(
    tmp_path: Path,
    limits: DocumentLimits,
    name: str,
    body: bytes,
    category: str,
) -> None:
    path = tmp_path / name
    path.write_bytes(body)
    with pytest.raises(DocumentParseError) as error:
        detect_document_type(path, filename=name, content_type=None, limits=limits)
    assert error.value.category == category


def test_conflicting_extension_and_mime_are_rejected(
    tmp_path: Path, limits: DocumentLimits
) -> None:
    path = tmp_path / "lecture.pdf"
    make_pdf(path, ["text"])
    with pytest.raises(DocumentParseError, match="type_mismatch"):
        detect_document_type(
            path,
            filename=path.name,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            limits=limits,
        )


def test_zip_traversal_and_ratio_limits_are_rejected(
    tmp_path: Path, limits: DocumentLimits
) -> None:
    traversal = tmp_path / "traversal.docx"
    with zipfile.ZipFile(traversal, "w") as archive:
        archive.writestr("word/document.xml", "<x/>")
        archive.writestr("../escape", "x")
    with pytest.raises(DocumentParseError, match="archive_unsafe_path"):
        detect_document_type(
            traversal,
            filename=traversal.name,
            content_type=None,
            limits=limits,
        )

    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 100_000)
    tight = limits.model_copy(update={"max_archive_ratio": 2})
    with pytest.raises(DocumentParseError, match="archive_ratio_limit"):
        detect_document_type(bomb, filename=bomb.name, content_type=None, limits=tight)


def test_character_and_unit_limits_are_enforced(tmp_path: Path, limits: DocumentLimits) -> None:
    path = tmp_path / "large.txt"
    path.write_text("x" * 100, encoding="utf-8")
    with pytest.raises(DocumentParseError, match="character_limit"):
        parse_document(path, "text", limits.model_copy(update={"max_characters": 10}))
    with pytest.raises(DocumentParseError, match="unit_limit"):
        parse_document(
            path,
            "text",
            limits.model_copy(update={"max_unit_characters": 10, "max_units": 2}),
        )
    with pytest.raises(DocumentParseError, match="source_size_limit"):
        parse_document(path, "text", limits.model_copy(update={"max_source_bytes": 10}))
