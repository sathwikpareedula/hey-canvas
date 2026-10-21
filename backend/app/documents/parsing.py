"""Bounded, local-only parsers for allowlisted academic document formats."""

import hashlib
import re
import unicodedata
import zipfile
from collections.abc import Iterable
from pathlib import Path, PurePosixPath

from docx import Document as OpenDocx
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader

from app.documents.models import DocumentLimits, DocumentType, ParsedUnit, ParseResult

PARSER_NAME = "uta-academic-document-parser"
PARSER_VERSION = "1"
CONTROL_NOISE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
BLANK_LINES = re.compile(r"\n{3,}")

EXTENSION_TYPES: dict[str, DocumentType] = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".txt": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".csv": "csv",
}
MIME_TYPES: dict[str, DocumentType] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text",
    "text/markdown": "markdown",
    "text/csv": "csv",
}


class DocumentParseError(Exception):
    """Sanitized parser failure suitable for persistence and CLI reporting."""

    def __init__(self, category: str) -> None:
        super().__init__(category)
        self.category = category


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFC", value.replace("\r\n", "\n").replace("\r", "\n"))
    value = CONTROL_NOISE.sub("", value)
    lines = [line.rstrip() for line in value.split("\n")]
    return BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


def supported_candidate(filename: str | None, content_type: str | None) -> bool:
    extension = Path(filename or "").suffix.lower()
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    return extension in EXTENSION_TYPES or mime in MIME_TYPES


def detect_document_type(
    path: Path,
    *,
    filename: str | None,
    content_type: str | None,
    limits: DocumentLimits,
) -> DocumentType:
    if path.stat().st_size > limits.max_source_bytes:
        raise DocumentParseError("source_size_limit")
    extension_type = EXTENSION_TYPES.get(Path(filename or path.name).suffix.lower())
    mime_type = MIME_TYPES.get((content_type or "").split(";", 1)[0].strip().lower())
    if extension_type is not None and mime_type is not None and extension_type != mime_type:
        raise DocumentParseError("type_mismatch")
    with path.open("rb") as source:
        signature = source.read(8)
    if signature.startswith(b"%PDF-"):
        actual: DocumentType = "pdf"
    elif signature.startswith(b"PK\x03\x04"):
        actual = _detect_office_archive(path, limits)
    elif signature.startswith(b"MZ") or signature.startswith(b"\x7fELF"):
        raise DocumentParseError("unsupported_executable")
    elif extension_type in {"text", "markdown", "csv"} or mime_type in {
        "text",
        "markdown",
        "csv",
    }:
        actual = extension_type or mime_type or "text"
    else:
        category = "type_mismatch" if extension_type or mime_type else "unsupported_type"
        raise DocumentParseError(category)
    declared = extension_type or mime_type
    if declared is not None and declared != actual:
        raise DocumentParseError("type_mismatch")
    return actual


def parse_document(path: Path, document_type: DocumentType, limits: DocumentLimits) -> ParseResult:
    if path.stat().st_size > limits.max_source_bytes:
        raise DocumentParseError("source_size_limit")
    try:
        if document_type == "pdf":
            units = _parse_pdf(path, limits)
            status = "needs_ocr" if sum(unit.character_count for unit in units) < 20 else "parsed"
        elif document_type == "pptx":
            _validate_archive(path, limits)
            units = _parse_pptx(path, limits)
            status = "parsed"
        elif document_type == "docx":
            _validate_archive(path, limits)
            units = _parse_docx(path, limits)
            status = "parsed"
        else:
            units = _parse_text(path, limits, document_type)
            status = "parsed"
    except DocumentParseError:
        raise
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise DocumentParseError("malformed_document") from exc
    except Exception as exc:
        raise DocumentParseError("parser_error") from exc
    total = sum(unit.character_count for unit in units)
    if total > limits.max_characters:
        raise DocumentParseError("character_limit")
    if len(units) > limits.max_units:
        raise DocumentParseError("unit_limit")
    return ParseResult(
        detected_type=document_type,
        status=status,
        parser_name=PARSER_NAME,
        parser_version=PARSER_VERSION,
        units=tuple(units),
        character_count=total,
    )


def _detect_office_archive(path: Path, limits: DocumentLimits) -> DocumentType:
    names = _validate_archive(path, limits)
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    raise DocumentParseError("unsupported_archive")


def _validate_archive(path: Path, limits: DocumentLimits) -> set[str]:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise DocumentParseError("malformed_archive") from exc
    if len(infos) > limits.max_archive_entries:
        raise DocumentParseError("archive_entry_limit")
    total = 0
    names: set[str] = set()
    for info in infos:
        normalized = info.filename.replace("\\", "/")
        parts = PurePosixPath(normalized).parts
        if normalized.startswith("/") or ".." in parts:
            raise DocumentParseError("archive_unsafe_path")
        if info.flag_bits & 0x1:
            raise DocumentParseError("encrypted_archive")
        total += info.file_size
        if total > limits.max_archive_uncompressed_bytes:
            raise DocumentParseError("archive_size_limit")
        if info.file_size and (
            info.compress_size == 0
            or info.file_size / info.compress_size > limits.max_archive_ratio
        ):
            raise DocumentParseError("archive_ratio_limit")
        names.add(normalized)
    return names


def _unit(
    *,
    unit_type: str,
    unit_index: int,
    text: str,
    page_number: int | None = None,
    slide_number: int | None = None,
    heading: str | None = None,
) -> ParsedUnit:
    normalized = normalize_text(text)
    return ParsedUnit(
        unit_type=unit_type,
        unit_index=unit_index,
        page_number=page_number,
        slide_number=slide_number,
        heading=normalize_text(heading)[:1000] if heading else None,
        text=normalized,
        text_fingerprint=hashlib.sha256(normalized.encode()).hexdigest(),
        character_count=len(normalized),
    )


def _bounded_units(
    blocks: Iterable[tuple[str, int | None, int | None, str | None, str]],
    limits: DocumentLimits,
) -> list[ParsedUnit]:
    result: list[ParsedUnit] = []
    for unit_type, page, slide, heading, text in blocks:
        normalized = normalize_text(text)
        chunks = _chunks(normalized, limits.max_unit_characters) or [""]
        for chunk in chunks:
            if len(result) >= limits.max_units:
                raise DocumentParseError("unit_limit")
            result.append(
                _unit(
                    unit_type=unit_type,
                    unit_index=len(result),
                    page_number=page,
                    slide_number=slide,
                    heading=heading,
                    text=chunk,
                )
            )
            if sum(item.character_count for item in result) > limits.max_characters:
                raise DocumentParseError("character_limit")
    return result


def _chunks(text: str, maximum: int) -> list[str]:
    if not text:
        return []
    chunks: list[str] = []
    remaining = text
    while len(remaining) > maximum:
        split = remaining.rfind("\n", 0, maximum + 1)
        if split < maximum // 2:
            split = maximum
        chunks.append(remaining[:split].rstrip())
        remaining = remaining[split:].lstrip("\n")
    if remaining:
        chunks.append(remaining)
    return chunks


def _parse_pdf(path: Path, limits: DocumentLimits) -> list[ParsedUnit]:
    reader = PdfReader(path, strict=False)
    if len(reader.pages) > limits.max_units:
        raise DocumentParseError("unit_limit")
    blocks: list[tuple[str, int | None, int | None, str | None, str]] = []
    for number, page in enumerate(reader.pages, start=1):
        contents = page.get_contents()
        if (
            contents is not None
            and len(contents.get_data()) > limits.max_archive_uncompressed_bytes
        ):
            raise DocumentParseError("pdf_stream_limit")
        blocks.append(("page", number, None, None, page.extract_text() or ""))
    return _bounded_units(blocks, limits)


def _parse_pptx(path: Path, limits: DocumentLimits) -> list[ParsedUnit]:
    presentation = Presentation(str(path))
    blocks: list[tuple[str, int | None, int | None, str | None, str]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        title = slide.shapes.title.text if slide.shapes.title is not None else None
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = getattr(shape, "text", "")
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        blocks.append(("slide", None, number, title, "\n".join(parts)))
    return _bounded_units(blocks, limits)


def _parse_docx(path: Path, limits: DocumentLimits) -> list[ParsedUnit]:
    document = OpenDocx(str(path))
    sections: list[tuple[str | None, list[str]]] = [(None, [])]
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text
            if block.style is not None and block.style.name.startswith("Heading") and text.strip():
                sections.append((text, []))
            elif text.strip():
                sections[-1][1].append(text)
        elif isinstance(block, Table):
            sections[-1][1].extend(
                " | ".join(cell.text for cell in row.cells) for row in block.rows
            )
    blocks = [
        ("section", None, None, heading, "\n".join(parts))
        for heading, parts in sections
        if heading or parts
    ]
    return _bounded_units(blocks, limits)


def _parse_text(path: Path, limits: DocumentLimits, kind: DocumentType) -> list[ParsedUnit]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    return _bounded_units([(kind, None, None, None, text)], limits)
